from django.shortcuts import render, redirect, get_object_or_404
from django.http import HttpResponse
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from rest_framework import viewsets, permissions
from django_filters.rest_framework import DjangoFilterBackend
from accounts.decorators import puede_editar, rol_requerido
from .models import Programa
from .serializers import ProgramaSerializer
from .forms import ProgramaForm


class ProgramaViewSet(viewsets.ModelViewSet):
    queryset         = Programa.objects.select_related('sede__facultad__universidad')
    serializer_class = ProgramaSerializer
    filter_backends  = [DjangoFilterBackend]
    filterset_fields = ['sede', 'estado', 'area', 'activo']

    def get_permissions(self):
        if self.action in ['list', 'retrieve']:
            return [permissions.IsAuthenticated()]
        return [permissions.IsAdminUser()]


@login_required
def lista_programas(request):
    programas_activos = Programa.objects.select_related(
        'sede__facultad__universidad'
    ).filter(activo=True).order_by(
        'sede__facultad__universidad__sigla', 'nombre'
    )

    programas_inactivos = Programa.objects.select_related(
        'sede__facultad__universidad'
    ).filter(activo=False).order_by(
        'sede__facultad__universidad__sigla', 'nombre'
    )

    todos = list(programas_activos)
    ya_carrera = sum(
        1 for p in todos
        if p.estado == 'IMPLEMENTACION' and (
            'RAN' in (p.observaciones or '') or
            'UAGRM' in (p.observaciones or '') or
            'UAJMS' in (p.observaciones or '')
        )
    )

    conteo = {
        'formulacion':    sum(1 for p in todos if p.estado == 'FORMULACION'),
        'aprobacion':     sum(1 for p in todos if p.estado == 'APROBACION'),
        'implementacion': sum(1 for p in todos if p.estado == 'IMPLEMENTACION'),
        'suspendido':     sum(1 for p in todos if p.estado == 'SUSPENDIDO'),
        'ya_carrera':     ya_carrera,
    }

    return render(request, 'programas/lista.html', {
        'programas':          programas_activos,
        'programas_todos':    todos,
        'programas_inactivos': programas_inactivos,
        'conteo':             conteo,
    })


@login_required
@puede_editar
def crear_programa(request):
    form = ProgramaForm(request.POST or None)
    if request.method == 'POST' and form.is_valid():
        form.save()
        messages.success(request, '✅ Programa creado correctamente.')
        return redirect('lista-programas')
    return render(request, 'programas/form_programa.html',
                  {'form': form, 'titulo': 'Nuevo Programa', 'accion': 'Crear'})


@login_required
@puede_editar
def editar_programa(request, pk):
    programa = get_object_or_404(Programa, pk=pk)
    form     = ProgramaForm(request.POST or None, instance=programa)
    if request.method == 'POST' and form.is_valid():
        form.save()
        messages.success(request, '✅ Programa actualizado.')
        return redirect('lista-programas')
    return render(request, 'programas/form_programa.html',
                  {'form': form, 'titulo': f'Editar — {programa.nombre}',
                   'accion': 'Guardar cambios', 'objeto': programa})


@login_required
@rol_requerido('SUPERADMIN', 'ADMIN_DSA')
def eliminar_programa(request, pk):
    programa = get_object_or_404(Programa, pk=pk)
    if request.method == 'POST':
        nombre = str(programa)
        programa.delete()
        messages.warning(request, f'Programa {nombre} eliminado.')
        return redirect('lista-programas')
    return render(request, 'programas/confirmar_eliminar.html',
                  {'objeto': programa, 'tipo': 'Programa'})

@login_required
def exportar_programas_excel(request):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from django.utils import timezone as tz

    wb = openpyxl.Workbook()

    fill_h  = PatternFill('solid', fgColor='0F1F3D')
    fill_sub = PatternFill('solid', fgColor='1A3560')
    fill_par = PatternFill('solid', fgColor='EEF2F8')
    fill_ina = PatternFill('solid', fgColor='E8E8E8')
    font_h  = Font(bold=True, color='FFFFFF', size=10)
    centro  = Alignment(horizontal='center', vertical='center', wrap_text=True)
    izq     = Alignment(horizontal='left',   vertical='center', wrap_text=True)
    borde   = Border(
        left=Side(style='thin',   color='CCCCCC'),
        right=Side(style='thin',  color='CCCCCC'),
        top=Side(style='thin',    color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC'),
    )

    def cabecera(ws, fila, cols, fill=fill_h):
        for col, txt in enumerate(cols, 1):
            c = ws.cell(row=fila, column=col, value=txt)
            c.fill = fill; c.font = font_h
            c.alignment = centro; c.border = borde

    def fila_datos(ws, fila, vals, par=False, inactivo=False):
        for col, val in enumerate(vals, 1):
            c = ws.cell(row=fila, column=col, value=val)
            c.alignment = izq; c.border = borde
            if inactivo:
                c.fill = fill_ina
                c.font = Font(color='888888', italic=True)
            elif par:
                c.fill = fill_par

    # ── Hoja 1: Resumen ──────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = '1. Resumen Programas'
    ws1.sheet_view.showGridLines = False

    ws1.merge_cells('A1:E1')
    c = ws1['A1']
    c.value     = 'SIGCU Bolivia — Resumen de Programas Académicos'
    c.font      = Font(bold=True, size=13, color='0F1F3D')
    c.alignment = centro
    ws1.row_dimensions[1].height = 24

    ws1.merge_cells('A2:E2')
    c = ws1['A2']
    c.value     = (f'Universidad Autónoma Tomás Frías — DSA | '
                   f'Generado: {tz.now().strftime("%d/%m/%Y %H:%M")}')
    c.font      = Font(italic=True, size=9, color='666666')
    c.alignment = centro
    ws1.row_dimensions[3].height = 8

    from programas.models import Programa
    programas = Programa.objects.select_related('sede__facultad__universidad')

    resumen = [
        ('Total programas',         programas.count()),
        ('Activos',                 programas.filter(activo=True).count()),
        ('Inactivos',               programas.filter(activo=False).count()),
        ('En Formulación',          programas.filter(estado='FORMULACION').count()),
        ('En Aprobación',           programas.filter(estado='APROBACION').count()),
        ('En Implementación',       programas.filter(estado='IMPLEMENTACION').count()),
        ('Suspendidos',             programas.filter(estado='SUSPENDIDO').count()),
    ]
    cabecera(ws1, 4, ['INDICADOR', 'CANTIDAD', '', '', ''])
    for i, (label, val) in enumerate(resumen, 5):
        fila_datos(ws1, i, [label, val, '', '', ''], par=(i % 2 == 0))

    ws1.column_dimensions['A'].width = 30
    ws1.column_dimensions['B'].width = 14

    # ── Hoja 2: Detalle completo ──────────────────────────────────────
    ws2 = wb.create_sheet('2. Detalle Programas')
    ws2.sheet_view.showGridLines = False

    ws2.merge_cells('A1:M1')
    c = ws2['A1']
    c.value     = 'SIGCU Bolivia — Detalle Completo de Programas'
    c.font      = Font(bold=True, size=12, color='0F1F3D')
    c.alignment = centro
    ws2.row_dimensions[1].height = 22

    cols2 = [
        'N°', 'Universidad', 'Facultad', 'Sede', 'Ciudad',
        'Programa', 'Grado previsto', 'Estado programa',
        'Responsable', 'Fecha inicio', 'Fecha aprob. prevista',
        'Observaciones', 'Activo'
    ]
    cabecera(ws2, 2, cols2)
    ws2.row_dimensions[2].height = 28

    todos = Programa.objects.select_related(
        'sede__facultad__universidad'
    ).order_by(
        'sede__facultad__universidad__sigla',
        'sede__facultad__nombre', 'nombre'
    )

    for i, p in enumerate(todos, 1):
        inactivo = not p.activo
        fila_datos(ws2, i + 2, [
            i,
            p.sede.facultad.universidad.sigla,
            p.sede.facultad.nombre,
            p.sede.nombre,
            p.sede.ciudad,
            p.nombre,
            p.get_grado_previsto_display() if p.grado_previsto else '—',
            p.get_estado_display(),
            p.responsable or '—',
            p.fecha_inicio.strftime('%d/%m/%Y') if p.fecha_inicio else '—',
            p.fecha_prevista_aprobacion.strftime('%d/%m/%Y')
                if p.fecha_prevista_aprobacion else '—',
            p.observaciones or '—',
            'Sí' if p.activo else 'No',
        ], par=(i % 2 == 0), inactivo=inactivo)

    anchos2 = [5, 12, 38, 32, 14, 42, 16, 18, 22, 12, 18, 30, 8]
    for col, ancho in enumerate(anchos2, 1):
        ws2.column_dimensions[get_column_letter(col)].width = ancho
    ws2.freeze_panes = 'A3'

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=programas_sigcu.xlsx'
    wb.save(response)
    return response